"""Bounded, preview-first migration into native workforce and accounting records."""
import base64
import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
import io
import json
import re
import secrets
import zipfile

import frappe
from frappe.utils import getdate, today

from plus_point.finance import as_decimal, daily_amounts, day_units, money
from plus_point.setup_flow import linked, require_customer

SCHEMAS = {
    'employees': ['external_id', 'employee', 'first_name', 'last_name', 'gender', 'date_of_birth', 'date_of_joining', 'cell_number', 'status'],
    'work': ['external_id', 'assignment', 'work_date', 'units'],
    'milestones': ['external_id', 'contract_terms', 'title', 'date', 'amount'],
    'expenses': ['external_id', 'date', 'project', 'supplier', 'item_code', 'amount', 'description'],
    'invoices': ['external_id', 'date', 'project', 'customer', 'item_code', 'amount', 'source_type', 'source_name', 'description'],
    'receipts': ['external_id', 'date', 'invoice', 'amount', 'bank_account'],
    'worker_payments': ['external_id', 'date', 'project', 'employee', 'amount', 'bank_account'],
}
LEGACY_EMPLOYEE_COLUMNS = list(SCHEMAS['employees'])
EMPLOYEE_OPTIONAL_COLUMNS = ['middle_name', 'employee_number', 'personal_email', 'current_address',
    'permanent_address', 'person_to_be_contacted', 'emergency_phone_number', 'relation', 'passport_number', 'valid_upto']
SCHEMAS['employees'] += EMPLOYEE_OPTIONAL_COLUMNS
TARGETS = dict(employees='Employee', work='PP Work Entry', milestones='PP Milestone',
    expenses='Purchase Invoice', invoices='Sales Invoice', receipts='Payment Entry', worker_payments='PP Wage Settlement')
MAX_ROWS = 200
MAX_BYTES = 2 * 1024 * 1024
from plus_point.event_imports import SCHEMAS as EVENT_SCHEMAS, TARGETS as EVENT_TARGETS, PROFILES as EVENT_PROFILES
SCHEMAS.update(EVENT_SCHEMAS)
TARGETS.update(EVENT_TARGETS)


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def record_key(company, kind, external_id):
    return digest([company, kind, external_id])


def guard(kind, company):
    require_customer()
    if kind not in SCHEMAS:
        frappe.throw('Choose a supported import template')
    doc = linked('Company', company)
    frappe.has_permission(TARGETS[kind], 'create', throw=True)
    if kind not in ('employees','work','milestones') and kind not in EVENT_SCHEMAS:
        frappe.has_permission('Payment Entry' if kind in ('receipts', 'worker_payments') else TARGETS[kind], 'submit', throw=True)
    if kind=='stock_balances':frappe.has_permission('Stock Reconciliation','submit',throw=True)
    return doc


def parse_file(filename, content, kind):
    if not isinstance(filename, str) or not isinstance(content, str) or len(content) > MAX_BYTES * 1.4:
        frappe.throw('Upload a CSV or XLSX file smaller than 2 MB')
    try:
        data = base64.b64decode(content, validate=True)
        if len(data) > MAX_BYTES:
            raise ValueError()
        if filename.lower().endswith('.csv'):
            rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig')), strict=True))
        elif filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 10 * MAX_BYTES:
                    frappe.throw('This workbook is too large. Export the import sheet as CSV.')
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                if len(book.worksheets) != 1:
                    frappe.throw('Upload one worksheet only, or export the required sheet as CSV')
                sheet = book.worksheets[0]
                rows = []
                for cells in sheet.iter_rows(max_row=MAX_ROWS+2, max_col=len(SCHEMAS[kind])+1):
                    if any(cell.data_type == 'f' for cell in cells):
                        frappe.throw('Replace spreadsheet formulas with their values before uploading')
                    rows.append([cell.value if cell.value is not None else '' for cell in cells])
                if sheet.max_row and sheet.max_row > MAX_ROWS+1:
                    frappe.throw('Split this sheet into files of at most 200 rows')
            finally:
                book.close()
        else:
            frappe.throw('Use a UTF-8 CSV or XLSX file. Google Sheets: File → Download → CSV or Microsoft Excel.')
    except (ValueError, UnicodeError, csv.Error, zipfile.BadZipFile):
        frappe.throw('The file could not be read. Save it as UTF-8 CSV or XLSX and try again.')
    while rows and not any(str(value).strip() for value in rows[-1]):
        rows.pop()
    if not rows:
        frappe.throw('The file is empty. Download a template and add your records.')
    headers = [str(value).strip() for value in rows.pop(0)]
    while headers and not headers[-1]:
        headers.pop()
    if headers != SCHEMAS[kind] and not (kind == 'employees' and headers == LEGACY_EMPLOYEE_COLUMNS):
        frappe.throw('Column headers must match the selected template in the same order. Download a fresh template.')
    if not rows or len(rows) > MAX_ROWS:
        frappe.throw('Add between 1 and 200 data rows. Blank rows between records are not allowed.')
    result = []
    for number, values in enumerate(rows, 2):
        if len(values) > len(headers) and any(str(v).strip() for v in values[len(headers):]):
            frappe.throw(f'Row {number}: unexpected extra columns')
        values = list(values[:len(headers)]) + [''] * max(0, len(headers)-len(values))
        def text(value):
            return value.strftime('%Y-%m-%d') if isinstance(value, (datetime, date)) else str(value).strip()
        row = dict(zip(headers, map(text, values)))
        if any(len(v) > 2000 for v in row.values()):
            frappe.throw(f'Row {number}: a cell exceeds 2,000 characters')
        result.append(row)
    return result


def required(row, field):
    if not row.get(field):
        frappe.throw(f'Enter {field}')
    return row[field]


def row_date(row, field, historical=True):
    value = required(row, field)
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        frappe.throw(f'{field} must use YYYY-MM-DD')
    parsed = getdate(value)
    if historical and parsed > getdate(today()):
        frappe.throw(f'{field} cannot be in the future for a historical import')
    return parsed


def amount(row):
    value = as_decimal(required(row, 'amount'))
    if value <= 0 or value != money(value):
        frappe.throw('amount must be positive, with at most two decimal places and no currency symbols')
    return value


def resolve(doctype, value, company, import_kind=None):
    if value.startswith('import:') and import_kind:
        found = frappe.db.get_value('PP Import Record', record_key(company, import_kind, value[7:]), 'document_name')
        if not found:
            frappe.throw('Referenced import ID was not found. Import the earlier records first.')
        value = found
    return linked(doctype, value)


def employee_target(company, row):
    """A name in the old employee column is a label, not a required database ID."""
    value = row.get('employee', '')
    if not value:return None
    if frappe.db.exists('Employee', value):
        linked('Employee', value, company).check_permission('write')
        return value
    matches = frappe.get_list('Employee', filters={'company':company, 'employee_name':value},
        fields=['name', 'date_of_birth'], limit_page_length=3)
    if len(matches) > 1:
        frappe.throw('More than one worker has this name. Use the Employee ID from reference IDs to select the correct profile.')
    if matches:
        if getdate(matches[0].date_of_birth) != getdate(row.get('date_of_birth')):
            frappe.throw('A worker with this name has a different date of birth. Use their Employee ID to update them, or leave employee blank to add a different person.')
        linked('Employee', matches[0].name, company).check_permission('write')
        return matches[0].name
    return None


def validate_row(kind, company, row):
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._-]{0,79}', required(row, 'external_id')):
        frappe.throw('external_id must be 1–80 letters, digits, dots, underscores or hyphens')
    for key, value in row.items():
        if value[:1] in ('=', '@', '+') and not (key in ('cell_number','emergency_phone_number') and re.fullmatch(r'\+\d[\d ()-]*', value)):
            frappe.throw(f'{key}: replace formulas with plain values')
    if kind in EVENT_SCHEMAS:
        from plus_point.event_imports import validate
        return validate(kind,company,row)
    if kind == 'employees':
        required(row, 'first_name')
        linked('Gender', required(row, 'gender'))
        dob, joining = row_date(row, 'date_of_birth'), row_date(row, 'date_of_joining')
        if dob >= joining:
            frappe.throw('date_of_birth must be before date_of_joining')
        if row['status'] not in ('Active', 'Inactive'):
            frappe.throw('status must be Active or Inactive')
        if row.get('personal_email'):
            frappe.utils.validate_email_address(row['personal_email'], throw=True)
        if row.get('valid_upto'):row_date(row, 'valid_upto', historical=False)
        employee_target(company, row)
        return dict(label=f"{row['first_name']} {row['last_name']}".strip(), amount=None)
    if kind == 'work':
        assignment = linked('PP Assignment', required(row, 'assignment'))
        linked('Project', assignment.project, company)
        day, units = row_date(row, 'work_date'), day_units(required(row, 'units'))
        if day < getdate(assignment.start_date) or (assignment.end_date and day > getdate(assignment.end_date)):
            frappe.throw('work_date must fall within the assignment dates')
        return dict(label=assignment.employee, amount=str(daily_amounts(assignment.wage_rate, 0, units)[0]), project=assignment.project)
    value = amount(row)
    day = row_date(row, 'date')
    if kind == 'milestones':
        terms = linked('PP Contract Terms', required(row, 'contract_terms'))
        linked('Project', terms.project, company)
        if terms.billing_model != 'Fixed' or terms.status != 'Active':
            frappe.throw('Choose active Fixed billing terms')
        required(row, 'title')
        return dict(label=row['title'], amount=str(value), project=terms.project)
    if kind == 'receipts':
        invoice = resolve('Sales Invoice', required(row, 'invoice'), company, 'invoices')
        if invoice.company != company or invoice.docstatus != 1 or invoice.is_return:
            frappe.throw('Choose a submitted client invoice in the selected company')
        if invoice.currency != frappe.db.get_value('Company', company, 'default_currency'):
            frappe.throw('Historical imports require the company currency')
        bank = linked('Account', required(row, 'bank_account'), company)
        if bank.is_group or bank.account_type not in ('Bank', 'Cash') or bank.account_currency != invoice.currency:
            frappe.throw('Choose a cash/bank ledger in the company currency')
        if day < getdate(invoice.posting_date):
            frappe.throw('Receipt date cannot be before invoice date')
        return dict(label=invoice.name, amount=str(value), project=invoice.project)
    project = linked('Project', required(row, 'project'), company)
    if kind == 'worker_payments':
        employee = resolve('Employee', required(row, 'employee'), company, 'employees')
        if employee.company != company:
            frappe.throw('Worker must belong to the selected company')
        bank = linked('Account', required(row, 'bank_account'), company)
        if bank.is_group or bank.account_type not in ('Bank', 'Cash') or bank.account_currency != frappe.db.get_value('Company', company, 'default_currency'):
            frappe.throw('Choose a cash/bank ledger in the company currency')
        return dict(label=employee.name, amount=str(value), project=project.name)
    party_type = 'Supplier' if kind == 'expenses' else 'Customer'
    party = linked(party_type, required(row, party_type.lower()))
    item = linked('Item', required(row, 'item_code'))
    if item.is_stock_item or item.disabled:
        frappe.throw('Choose an enabled non-stock service item')
    if kind == 'invoices':
        if project.customer != party.name:
            frappe.throw('Client invoice must match the project client')
        if bool(row['source_type']) != bool(row['source_name']):
            frappe.throw('Fill both source_type and source_name, or leave both blank for a legacy invoice')
        if row['source_type']:
            if row['source_type'] not in ('PP Work Entry', 'PP Milestone'):
                frappe.throw('source_type must be PP Work Entry or PP Milestone')
            source = resolve(row['source_type'], row['source_name'], company,
                             'work' if row['source_type'] == 'PP Work Entry' else 'milestones')
            if source.status != 'Approved' or source.sales_invoice:
                frappe.throw('Invoice source must be approved and unbilled')
            terms = linked('PP Contract Terms', source.contract_terms)
            expected = source.billable_amount if row['source_type'] == 'PP Work Entry' else source.amount
            if terms.project != project.name or value != money(expected):
                frappe.throw('Invoice amount and project must match the approved source')
    return dict(label=party.name, amount=str(value), project=project.name)


def prepare(kind, company, rows):
    results, seen, targets = [], set(), set()
    for number, row in enumerate(rows, 2):
        result = dict(row=number, external_id=row.get('external_id', ''), action='Create', doctype=TARGETS[kind])
        message_count = len(frappe.message_log)
        try:
            if row['external_id'] in seen:
                frappe.throw('external_id appears twice in this file')
            seen.add(row['external_id'])
            previous = frappe.db.get_value('PP Import Record', record_key(company, kind, row['external_id']),
                ['document_name', 'content_hash', 'document_type'], as_dict=True)
            if previous:
                if not frappe.db.exists(previous.document_type, previous.document_name):
                    frappe.throw('The earlier imported record is missing. Ask your administrator to review its import history.')
                target = linked(previous.document_type, previous.document_name)
                if target.docstatus == 2 or (kind == 'work' and target.status != 'Approved'):
                    frappe.throw('The earlier financial record was cancelled or changed. Use a new external ID for its reviewed replacement.')
                if kind in EVENT_PROFILES:
                    from plus_point.event_imports import target as catalog_target
                    current=catalog_target(kind,row)
                    if current and current!=previous.document_name:frappe.throw('external_id is already linked to a different catalog record')
                    target.check_permission('write')
                    if kind=='vehicles' and target.company!=company:frappe.throw('Vehicle must belong to the selected company')
                    if kind=='event_sites':linked('Project',target.project,company)
                    result['action']='Update' if previous.content_hash!=digest(row) else 'Skip'
                elif kind == 'employees':
                    current = employee_target(company, row)
                    if current and current != previous.document_name:
                        frappe.throw('external_id is already linked to a different employee')
                    linked('Employee', previous.document_name, company).check_permission('write')
                    result['action'] = 'Update' if previous.content_hash != digest(row) else 'Skip'
                elif previous.content_hash != digest(row):
                    frappe.throw('This external_id was already posted with different data. Correct/cancel the original record before importing a replacement with a new ID.')
                else:
                    result['action'] = 'Skip'
                result['document'] = previous.document_name
            elif kind == 'employees':
                current = employee_target(company, row)
                if current:result.update(action='Update', document=current)
            elif kind in EVENT_PROFILES:
                from plus_point.event_imports import target as catalog_target
                current=catalog_target(kind,row)
                if current and frappe.db.exists(TARGETS[kind],current):result.update(action='Update',document=current)
            # A posted source may now be reserved by its imported invoice. An
            # identical financial replay reads the existing receipt, not the
            # create-only source rules. Permissions/company remain checked.
            if kind != 'employees' and result['action'] == 'Skip':
                result.update(label=target.name, amount=None)
                results.append(result)
                continue
            details = validate_row(kind, company, row)
            if kind in EVENT_PROFILES or kind=='stock_balances':
                from plus_point.event_imports import target as catalog_target
                target_key=(kind,catalog_target(kind,row) or row.get('customer_name')) if kind!='stock_balances' else (kind,row['item_code'],row['warehouse'])
                if target_key in targets:frappe.throw('The same catalog record or warehouse/item balance is targeted twice in this file')
                targets.add(target_key)
            if kind == 'employees' and result.get('document'):
                if result['document'] in targets:
                    frappe.throw('The same employee is targeted twice in this file')
                targets.add(result['document'])
            elif kind == 'employees':
                identity=tuple(row.get(k,'').casefold() for k in ('first_name','middle_name','last_name','date_of_birth'))
                if identity in targets:frappe.throw('The same new worker appears twice in this file. Keep one row per person.')
                targets.add(identity)
            result.update(details)
        except (frappe.ValidationError, frappe.PermissionError, frappe.DoesNotExistError, ValueError) as exc:
            result.update(action='Fix', error=frappe.utils.strip_html(str(exc)))
            del frappe.message_log[message_count:]
        results.append(result)
    return results


@frappe.whitelist()
def template(kind, company):
    guard(kind, company)
    output = io.StringIO(newline='')
    csv.writer(output).writerow(SCHEMAS[kind])
    frappe.response.update(type='download', filename=f'plus-point-{kind}.csv', filecontent=output.getvalue().encode('utf-8-sig'))


@frappe.whitelist()
def references(company):
    require_customer()
    linked('Company', company)
    projects = frappe.get_list('Project', filters={'company':company}, fields=['name','project_name'], limit_page_length=0)
    names = [p.name for p in projects]
    result = []
    def add(doctype, fields, filters=None):
        for row in frappe.get_list(doctype, fields=fields, filters=filters or {}, limit_page_length=1000):
            result.append([doctype, row[fields[0]], row.get(fields[1], '') if len(fields)>1 else ''])
    result.extend(['Project',p.name,p.project_name] for p in projects)
    add('Employee',['name','employee_name'],{'company':company})
    for doctype, fields in [('PP Assignment',['name','employee']),('PP Contract Terms',['name','billing_model'])]:
        if names:
            add(doctype,fields,{'project':['in',names]})
    add('Customer',['name','customer_name'])
    add('Supplier',['name','supplier_name'])
    add('Item',['name','item_name'],{'is_stock_item':0,'disabled':0})
    add('Item',['name','item_name'],{'is_stock_item':1,'disabled':0})
    add('Warehouse',['name','warehouse_name'],{'company':company,'is_group':0,'disabled':0})
    add('Vehicle',['name','model'],{'company':company})
    if names:add('PP Site',['name','site_name'],{'project':['in',names]})
    add('Account',['name','account_type'],{'company':company,'is_group':0,'account_type':['in',['Bank','Cash']]})
    add('Sales Invoice',['name','customer'],{'company':company,'docstatus':1,'outstanding_amount':['>',0]})
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(['record_type','id','label'])
    for row in result:
        writer.writerow([("'"+str(value)) if str(value)[:1] in ('=','+','-','@') else value for value in row])
    frappe.response.update(type='download',filename='plus-point-reference-ids.csv',filecontent=output.getvalue().encode('utf-8-sig'))


def spending_states(kind, results):
    if kind not in ('work', 'expenses'):
        return []
    totals = {}
    for row in results:
        if row.get('project') and row['action'] not in ('Skip','Fix'):
            totals[row['project']] = totals.get(row['project'],Decimal(0)) + as_decimal(row['amount'])
    from plus_point.budget import status
    return [status(project, proposed=value) for project,value in sorted(totals.items())]


@frappe.whitelist(methods=['POST'])
def preview(kind, company, filename, content):
    company_doc = guard(kind, company)
    rows = parse_file(filename, content, kind)
    results = prepare(kind, company, rows)
    budgets = spending_states(kind, results)
    token = secrets.token_urlsafe(32)
    valid = not any(row.get('error') for row in results) and not any(row['blocked'] for row in budgets)
    if valid:
        frappe.cache.set_value('pp-import:' + token, dict(user=frappe.session.user, company=company, kind=kind, rows=rows), expires_in_sec=1200)
    return dict(rows=results, budgets=budgets, token=token if valid else None, currency=company_doc.default_currency,
                valid=valid, total=str(money(sum((as_decimal(row.get('amount') or 0) for row in results if row['action'] != 'Skip'), Decimal(0)))))


def create_financial(kind, company, row):
    if kind == 'work':
        doc = frappe.get_doc(dict(doctype='PP Work Entry', assignment=row['assignment'],
            work_date=row['work_date'], units=float(day_units(row['units'])), status='Draft',
            remarks='Historical import ' + row['external_id'])).insert()
        doc.status = 'Approved'
        doc.save()
        return doc
    if kind == 'milestones':
        return frappe.get_doc(dict(doctype='PP Milestone', contract_terms=row['contract_terms'],
            milestone_title=row['title'], due_date=row['date'], amount=float(amount(row)), status='Approved')).insert()
    if kind in ('expenses', 'invoices'):
        expense = kind == 'expenses'
        item = dict(item_code=row['item_code'], qty=1, rate=float(amount(row)), project=row['project'], cost_center=company.cost_center)
        if row['description']:
            item['description'] = row['description']
        item['expense_account' if expense else 'income_account'] = company.default_expense_account if expense else company.default_income_account
        if not expense and row['source_type']:
            source = resolve(row['source_type'], row['source_name'], company.name,
                             'work' if row['source_type'] == 'PP Work Entry' else 'milestones')
            item.update(pp_source_type=row['source_type'], pp_source_name=source.name)
        doc = frappe.get_doc(dict(doctype=TARGETS[kind], company=company.name, project=row['project'],
            posting_date=row['date'], set_posting_time=1, due_date=row['date'], currency=company.default_currency,
            conversion_rate=1, taxes=[], items=[item], **({'supplier':row['supplier'], 'bill_no':row['external_id'], 'bill_date':row['date']} if expense else {'customer':row['customer']})))
        doc.insert()
        if money(doc.grand_total) != amount(row) or doc.taxes:
            frappe.throw('This invoice added taxes or charges. Use the native invoice form to import its tax breakdown.')
        doc.submit()
        return doc
    if kind == 'receipts':
        from erpnext.accounts.doctype.payment_entry.payment_entry import get_payment_entry
        invoice = resolve('Sales Invoice', row['invoice'], company.name, 'invoices')
        invoice = frappe.get_doc('Sales Invoice', invoice.name, for_update=True)
        if amount(row) > money(invoice.outstanding_amount):
            frappe.throw('Receipt exceeds the invoice outstanding amount')
        doc = get_payment_entry('Sales Invoice', invoice.name, party_amount=float(amount(row)),
                                bank_amount=float(amount(row)), bank_account=row['bank_account'], reference_date=row['date'])
        doc.update(dict(posting_date=row['date'], project=invoice.project, reference_no=row['external_id'], reference_date=row['date']))
        doc.insert()
        doc.submit()
        return doc
    employee = resolve('Employee', row['employee'], company.name, 'employees')
    doc = frappe.get_doc(dict(doctype='Payment Entry', company=company.name, payment_type='Pay',
        party_type='Employee', party=employee.name, project=row['project'], posting_date=row['date'],
        paid_from=row['bank_account'], paid_to=company.default_payable_account,
        paid_amount=float(amount(row)), received_amount=float(amount(row)), source_exchange_rate=1,
        target_exchange_rate=1, reference_no=row['external_id'], reference_date=row['date'])).insert()
    doc.submit()
    return frappe.get_doc(dict(doctype='PP Wage Settlement', employee=employee.name,
        project=row['project'], payment_entry=doc.name)).insert().submit()


@frappe.whitelist(methods=['POST'])
def apply(token):
    require_customer()
    payload = frappe.cache.get_value('pp-import:' + str(token))
    if not payload or payload['user'] != frappe.session.user:
        frappe.throw('This preview expired or belongs to another user. Upload the file again.')
    kind, company_name, rows = payload['kind'], payload['company'], payload['rows']
    company = guard(kind, company_name)
    # Serialize import IDs and take project locks in the same order as budget writers.
    frappe.db.get_value('Company', company_name, 'name', for_update=True)
    results = prepare(kind, company_name, rows)
    if any(row.get('error') for row in results):
        frappe.throw('Records changed since the preview. Upload again to review the current errors.')
    from plus_point.budget import lock_project
    for project in sorted({r['project'] for r in results if r.get('project')}):
        lock_project(project)
    frappe.db.savepoint('pp_import_apply')
    try:
        for row, result in zip(rows, results):
            if result['action'] == 'Skip':
                continue
            if kind in EVENT_SCHEMAS:
                from plus_point.event_imports import save as save_catalog
                doc=save_catalog(kind,company_name,row,result.get('document'))
            elif kind == 'employees':
                doc = linked('Employee', result['document'], company_name) if result.get('document') else frappe.new_doc('Employee')
                # Older templates omit optional columns; never erase fields they cannot show.
                doc.update({key: row[key] for key in SCHEMAS[kind] if key in row and key not in ('external_id', 'employee')})
                doc.company = company_name
                doc.create_user_automatically = 0
                doc.save()
            else:
                doc = create_financial(kind, company, row)
            key = record_key(company_name, kind, row['external_id'])
            # This internal receipt is immutable through Desk/REST. Normal target
            # document insert/save/submit permissions and hooks are never bypassed.
            if frappe.db.exists('PP Import Record', key):
                frappe.db.set_value('PP Import Record', key, 'content_hash', digest(row))
            else:
                frappe.flags.pp_import_record = True
                try:
                    frappe.get_doc(dict(doctype='PP Import Record', name=key, company=company_name, kind=kind,
                        external_id=row['external_id'], content_hash=digest(row), document_type=doc.doctype,
                        document_name=doc.name)).insert(ignore_permissions=True)
                finally:
                    frappe.flags.pp_import_record = False
            result['document'] = doc.name
    except Exception as exc:
        frappe.db.rollback(save_point='pp_import_apply')
        frappe.throw(f"Row {result['row']}: {frappe.utils.strip_html(str(exc))}. No rows from this import were saved.")
    return dict(rows=results, created=sum(r['action']=='Create' for r in results),
        updated=sum(r['action']=='Update' for r in results), skipped=sum(r['action']=='Skip' for r in results))
